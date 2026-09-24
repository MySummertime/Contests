"""Populate the supplied workbook and render publication-resolution figures."""
from pathlib import Path
import json
import os
import openpyxl
from PIL import Image,ImageDraw,ImageFont
from config import RESULTS,TEMPLATE

RELAY_HEADERS=['中继架次编号','中继无人机编号','能源组件编号','开始时刻（s）','悬停经度（°）','悬停纬度（°）','悬停海拔（m）','建链完成时刻（s）','服务结束时刻（s）','返回O01时刻（s）','架次能耗（kWh）']
COMM_HEADERS=['运输架次编号','通信阶段','开始时刻（s）','结束时刻（s）','保障方式','中继架次编号']

def save_workbook(relays,audit):
    if not audit['ok'] or not audit['continuous_certified']:
        raise ValueError('Workbook export requires all constraints and continuous-link certification')
    wb=openpyxl.load_workbook(TEMPLATE)
    rs=wb['Q3_中继架次']; cs=wb['Q3_通信保障']
    if [rs.cell(1,c).value for c in range(1,12)]!=RELAY_HEADERS: raise ValueError('Relay template headers changed')
    if [cs.cell(1,c).value for c in range(1,7)]!=COMM_HEADERS: raise ValueError('Communication template headers changed')
    for sheet,width in ((rs,11),(cs,6)):
        for row in sheet.iter_rows(min_row=2,max_col=width):
            for cell in row: cell.value=None
    for row,r in enumerate(relays,2):
        values=[r.id,r.drone,r.battery,r.start,r.site.lon,r.site.lat,r.site.alt,r.link_done,r.service_end,r.return_time,r.energy]
        for col,value in enumerate(values,1): rs.cell(row,col).value=value
    for row,values in enumerate(audit['communication_rows'],2):
        for col,value in enumerate(values,1): cs.cell(row,col).value=value
    wb.save(TEMPLATE)
    return TEMPLATE

def save_diagnostic_figure(ins,trips,relays,audit):
    font_file=find_times_new_roman()
    RESULTS.mkdir(exist_ok=True)
    image=Image.new('RGBA',(4800,2600),(255,255,255,0))
    draw=ImageDraw.Draw(image)
    title=ImageFont.truetype(str(font_file),92)
    label=ImageFont.truetype(str(font_file),64)
    draw.text((250,80),'Q3 Transport Schedule and Relay Windows',font=title,fill=(0,0,0,255))
    left,right,top,bottom=380,4550,350,2250
    draw.line((left,top,left,bottom),fill=(0,0,0,255),width=8)
    draw.line((left,bottom,right,bottom),fill=(0,0,0,255),width=8)
    end=max([t.return_time for t in trips]+[r.return_time for r in relays]+[1])
    def xx(t): return int(left+(right-left)*t/end)
    def yy(i): return int(top+(bottom-top)*(i-1)/max(len(trips)-1,1))
    for t in trips:
        y=yy(int(t.id[1:]))
        draw.line((xx(t.start),y,xx(t.return_time),y),fill=(59,91,146,255),width=7)
        draw.line((xx(t.takeoff),y,xx(t.delivery),y),fill=(29,158,137,255),width=13)
    for r in relays:
        draw.line((xx(r.link_done),bottom+60,xx(r.service_end),bottom+60),fill=(214,140,69,255),width=20)
    draw.text((2100,2380),'Time (s)',font=label,fill=(0,0,0,255))
    draw.text((35,1120),'Trip',font=label,fill=(0,0,0,255))
    path=RESULTS/'q3_schedule_diagnostic.png'
    image.save(path,dpi=(600,600))
    return path

def find_times_new_roman():
    candidates=[]
    explicit=os.environ.get('Q3_TIMES_FONT')
    if explicit: candidates.append(Path(explicit).expanduser())
    windir=os.environ.get('WINDIR')
    if windir: candidates.extend(Path(windir,'Fonts',name) for name in ('times.ttf','timesbd.ttf'))
    for folder in (Path('/Library/Fonts'),Path('/System/Library/Fonts/Supplemental'),Path.home()/'Library/Fonts'):
        candidates.extend(folder.glob('Times New Roman*.ttf'))
        candidates.extend(folder.glob('Times New Roman*.ttc'))
    for candidate in candidates:
        if candidate.is_file(): return candidate
    raise RuntimeError('Times New Roman not found; set Q3_TIMES_FONT to its font file')

def save_status(payload):
    RESULTS.mkdir(exist_ok=True)
    path=RESULTS/'q3_status.json'
    path.write_text(json.dumps(payload,ensure_ascii=False,indent=2,default=float),encoding='utf-8')
    return path
